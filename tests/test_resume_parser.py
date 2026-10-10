import io

import pytest
from docx import Document

from src.resume_parser import UnsupportedFileError, clean_text, extract_text


def test_extract_txt():
    assert extract_text(b"Jane Doe\n\n  Python  \n", "cv.txt") == "Jane Doe\nPython"


def test_extract_docx():
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Skills: Python, SQL")
    buf = io.BytesIO()
    doc.save(buf)
    assert extract_text(buf.getvalue(), "cv.DOCX") == "Jane Doe\nSkills: Python, SQL"


def test_unsupported_extension():
    with pytest.raises(UnsupportedFileError):
        extract_text(b"data", "cv.png")


def test_clean_text_removes_blank_lines():
    assert clean_text("a\n\n   \nb  ") == "a\nb"
